import csv
import io
import re
from pathlib import Path
from typing import List, Dict, Any, Tuple
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from docx import Document

def extract_text_from_file(file_path: Path) -> str:
    ext = file_path.suffix.lower()
    if ext in [".txt", ".md", ".json"]:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    elif ext == ".csv":
        text_lines = []
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            reader = csv.reader(f)
            for row in reader:
                text_lines.append(" | ".join(row))
        return "\n".join(text_lines)
    elif ext == ".docx":
        doc = Document(file_path)
        full_text = []
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text.strip())
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    full_text.append(" | ".join(cells))
        return "\n\n".join(full_text)
    elif ext in [".xlsx", ".xls"]:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        full_text = []
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            full_text.append(f"--- Sheet: {sheet_name} ---")
            for row in sheet.iter_rows(values_only=True):
                row_str = " | ".join([str(val) for val in row if val is not None])
                if row_str.strip():
                    full_text.append(row_str)
        return "\n".join(full_text)
    else:
        # Fallback reading
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()

def parse_rfp_excel(file_path: Path) -> List[Dict[str, Any]]:
    wb = openpyxl.load_workbook(file_path, data_only=True)
    questions = []
    
    # Check sheets
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
            
        # Detect header row (usually row 0, 1, or 2)
        header_idx = -1
        question_col = -1
        section_col = -1
        id_col = -1
        
        for r_idx, row in enumerate(rows[:5]):
            row_lower = [str(c).lower().strip() if c is not None else "" for c in row]
            for c_idx, val in enumerate(row_lower):
                if any(keyword in val for keyword in ["question", "requirement", "query", "item description", "security prompt", "criteria"]):
                    question_col = c_idx
                    header_idx = r_idx
                elif any(keyword in val for keyword in ["section", "category", "domain", "topic"]):
                    section_col = c_idx
                elif any(keyword in val for keyword in ["id", "ref", "no", "item #", "#"]):
                    id_col = c_idx

            if question_col != -1:
                break
                
        # If no explicit header matched, pick column with longest average text length
        if question_col == -1:
            col_lengths: Dict[int, List[int]] = {}
            for row in rows[1:10]:
                for c_idx, val in enumerate(row):
                    if val is not None:
                        col_lengths.setdefault(c_idx, []).append(len(str(val)))
            if col_lengths:
                question_col = max(col_lengths.keys(), key=lambda c: sum(col_lengths[c]) / len(col_lengths[c]))
                header_idx = 0

        # Extract question items
        for r_idx in range(header_idx + 1, len(rows)):
            row = rows[r_idx]
            if question_col < len(row) and row[question_col] is not None:
                q_text = str(row[question_col]).strip()
                if len(q_text) > 3 and not q_text.lower().startswith("sheet"):
                    sec_name = str(row[section_col]).strip() if section_col != -1 and section_col < len(row) and row[section_col] else sheet_name
                    q_id = str(row[id_col]).strip() if id_col != -1 and id_col < len(row) and row[id_col] else f"Q-{len(questions)+1}"
                    
                    questions.append({
                        "id": q_id,
                        "sheet": sheet_name,
                        "row_index": r_idx + 1,  # 1-indexed for openpyxl
                        "section": sec_name,
                        "question": q_text,
                        "question_col": question_col + 1
                    })
    return questions

def parse_rfp_csv(file_path: Path) -> List[Dict[str, Any]]:
    questions = []
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        reader = list(csv.reader(f))
        if not reader:
            return []
        
        header = reader[0]
        q_col = 0
        for i, col in enumerate(header):
            if any(k in col.lower() for k in ["question", "requirement", "query", "description", "item"]):
                q_col = i
                break
                
        for idx, row in enumerate(reader[1:], start=1):
            if q_col < len(row) and row[q_col].strip():
                questions.append({
                    "id": f"Q-{idx}",
                    "sheet": "Sheet1",
                    "row_index": idx + 1,
                    "section": row[0] if len(row) > 1 and q_col != 0 else "General",
                    "question": row[q_col].strip(),
                    "question_col": q_col + 1
                })
    return questions

def parse_rfp_docx(file_path: Path) -> List[Dict[str, Any]]:
    doc = Document(file_path)
    questions = []
    
    # Check tables first
    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            cells = [cell.text.strip() for cell in row.cells]
            if len(cells) >= 1:
                first_text = cells[0]
                # Look for questions or requirements
                if "?" in first_text or (len(first_text) > 15 and r_idx > 0):
                    questions.append({
                        "id": f"T{t_idx+1}-R{r_idx+1}",
                        "sheet": f"Table-{t_idx+1}",
                        "row_index": r_idx,
                        "section": f"Table {t_idx+1}",
                        "question": first_text,
                        "question_col": 1
                    })

    # If no tables with questions, check paragraphs
    if not questions:
        for p_idx, para in enumerate(doc.paragraphs):
            text = para.text.strip()
            if text.endswith("?") or re.match(r"^(\d+[\.\)]|[A-Z][\.\)])\s+", text):
                questions.append({
                    "id": f"P-{p_idx+1}",
                    "sheet": "Document",
                    "row_index": p_idx + 1,
                    "section": "General",
                    "question": text,
                    "question_col": 1
                })
    return questions

def export_filled_excel(original_path: Path, filled_items: List[Dict[str, Any]], export_path: Path) -> None:
    wb = openpyxl.load_workbook(original_path)
    
    # Group items by sheet
    sheet_items: Dict[str, List[Dict[str, Any]]] = {}
    for item in filled_items:
        sheet_items.setdefault(item.get("sheet", "Sheet1"), []).append(item)
        
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Professional navy
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    cell_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    
    for sheet_name, items in sheet_items.items():
        if sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
        else:
            sheet = wb.active
            
        max_col = sheet.max_column
        answer_col = max_col + 1
        conf_col = max_col + 2
        src_col = max_col + 3
        
        # Headers
        sheet.cell(row=1, column=answer_col, value="RFPilot AI Answer").fill = header_fill
        sheet.cell(row=1, column=answer_col).font = header_font
        sheet.cell(row=1, column=answer_col).alignment = Alignment(horizontal="center", vertical="center")
        
        sheet.cell(row=1, column=conf_col, value="Confidence").fill = header_fill
        sheet.cell(row=1, column=conf_col).font = header_font
        sheet.cell(row=1, column=conf_col).alignment = Alignment(horizontal="center", vertical="center")
        
        sheet.cell(row=1, column=src_col, value="Evidence / Source Docs").fill = header_fill
        sheet.cell(row=1, column=src_col).font = header_font
        sheet.cell(row=1, column=src_col).alignment = Alignment(horizontal="center", vertical="center")
        
        # Fill in rows
        for item in items:
            r = item.get("row_index", 2)
            ans = item.get("answer", "")
            conf = item.get("confidence", "High")
            sources = item.get("sources", "")
            
            c_ans = sheet.cell(row=r, column=answer_col, value=ans)
            c_ans.alignment = Alignment(wrap_text=True, vertical="top")
            c_ans.border = cell_border
            
            c_conf = sheet.cell(row=r, column=conf_col, value=conf)
            c_conf.alignment = Alignment(horizontal="center", vertical="top")
            c_conf.border = cell_border
            
            c_src = sheet.cell(row=r, column=src_col, value=sources)
            c_src.alignment = Alignment(wrap_text=True, vertical="top")
            c_src.border = cell_border

        # Adjust column widths nicely
        sheet.column_dimensions[openpyxl.utils.get_column_letter(answer_col)].width = 50
        sheet.column_dimensions[openpyxl.utils.get_column_letter(conf_col)].width = 16
        sheet.column_dimensions[openpyxl.utils.get_column_letter(src_col)].width = 30
        
    wb.save(export_path)

def export_filled_docx(original_path: Path, filled_items: List[Dict[str, Any]], export_path: Path) -> None:
    doc = Document(original_path) if original_path.exists() else Document()
    doc.add_heading("RFPilot - Answered Questionnaire", level=1)
    
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    hdr[0].text = "ID"
    hdr[1].text = "Question / Requirement"
    hdr[2].text = "AI Answer"
    hdr[3].text = "Confidence & Sources"
    
    for item in filled_items:
        row = table.add_row().cells
        row[0].text = str(item.get("id", ""))
        row[1].text = str(item.get("question", ""))
        row[2].text = str(item.get("answer", ""))
        row[3].text = f"Confidence: {item.get('confidence', 'High')}\nSources: {item.get('sources', '')}"
        
    doc.save(export_path)
